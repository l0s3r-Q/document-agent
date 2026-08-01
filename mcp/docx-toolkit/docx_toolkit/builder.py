"""builder.py —— 按 DocumentSpec 生成 .docx：字体/字号/行距/缩进/对齐/表格。"""

from __future__ import annotations

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .styles import BUILTIN_STYLES

_ALIGN = {
    "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
    "LEFT": WD_ALIGN_PARAGRAPH.LEFT,
    "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT,
    "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _set_run_font(run, font_name: str | None, size_pt: float | None, bold: bool | None):
    """设置 run 字体（中西文都设置，中文必须写 w:eastAsia）。"""
    if font_name:
        run.font.name = font_name
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), font_name)
    if size_pt:
        run.font.size = Pt(float(size_pt))
    if bold is not None:
        run.font.bold = bold


def _apply_paragraph_format(pf, style: dict):
    """按样式字典设置段落格式。"""
    align = style.get("align")
    if align and align in _ALIGN:
        pf.alignment = _ALIGN[align]

    rule = style.get("line_spacing_rule")
    if rule == "EXACTLY" and style.get("line_spacing_pt"):
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(float(style["line_spacing_pt"]))
    elif rule == "AT_LEAST" and style.get("line_spacing_pt"):
        pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        pf.line_spacing = Pt(float(style["line_spacing_pt"]))
    elif style.get("line_spacing_multiple"):
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = float(style["line_spacing_multiple"])

    if style.get("first_line_indent_pt"):
        pf.first_line_indent = Pt(float(style["first_line_indent_pt"]))
    elif style.get("first_line_indent_chars") and style.get("size_pt"):
        # 缩进 N 字符 ≈ N × 字号
        pf.first_line_indent = Pt(float(style["size_pt"]) * float(style["first_line_indent_chars"]))


def _resolve_style(doc_type: str, role: str, override: dict | None = None, spec_styles: dict | None = None) -> dict:
    """取预置样式，叠加 spec.styles（按 role 覆盖）与 section 级覆盖。"""
    base = BUILTIN_STYLES.get(doc_type, BUILTIN_STYLES["general"])["roles"].get(role, {})
    merged = dict(base)
    if spec_styles and role in spec_styles:
        merged.update(spec_styles[role])
    if override:
        merged.update(override)
    return merged


def build(spec: dict, output_path: str) -> dict:
    """按 DocumentSpec 生成 docx。

    spec 契约:
      {doc_type, title, author?, date?, page{top_cm,bottom_cm,left_cm,right_cm},
       styles?{role: {font_name,size_pt,bold,align,line_spacing_rule,...}} 按 role 覆盖预置排版,
       sections:[{type,text,font?,align?,bold?,indent_first_line?,rows?,items?}]}
    """
    doc = Document()
    doc_type = spec.get("doc_type", "general")

    # ── 页面设置 ──────────────────────────────────────────────
    sec = doc.sections[0]
    page = spec.get("page") or BUILTIN_STYLES.get(doc_type, BUILTIN_STYLES["general"])["page"]
    sec.top_margin = Cm(page.get("top_cm", 2.54))
    sec.bottom_margin = Cm(page.get("bottom_cm", 2.54))
    sec.left_margin = Cm(page.get("left_cm", 3.17))
    sec.right_margin = Cm(page.get("right_cm", 3.17))

    spec_styles = spec.get("styles") or {}

    def add_paragraph(sec_spec: dict, role: str):
        override = {}
        if sec_spec.get("font"):
            override.update(sec_spec["font"])
        if sec_spec.get("align"):
            override["align"] = sec_spec["align"]
        if sec_spec.get("bold") is not None:
            override["bold"] = sec_spec["bold"]
        style = _resolve_style(doc_type, role, override, spec_styles)

        p = doc.add_paragraph()
        # 标题设置 Word 内置 Heading 样式，确保大纲/结构可识别
        if role in ("heading1", "heading2", "heading3"):
            style_name = {"heading1": "Heading 1", "heading2": "Heading 2", "heading3": "Heading 3"}[role]
            try:
                p.style = doc.styles[style_name]
            except KeyError:
                pass
        _apply_paragraph_format(p.paragraph_format, style)
        text = sec_spec.get("text", "")
        if text:
            run = p.add_run(text)
            _set_run_font(
                run,
                style.get("font_name"),
                style.get("size_pt"),
                style.get("bold"),
            )
        return p

    # ── 标题 ──────────────────────────────────────────────────
    title = spec.get("title")
    if title:
        add_paragraph({"text": title}, "title")

    # ── 作者/日期行（可选）─────────────────────────────────────
    meta_line = "  ".join(x for x in (spec.get("author"), spec.get("date")) if x)
    if meta_line:
        add_paragraph({"text": meta_line, "align": "CENTER"}, "body")

    # ── 正文各节 ──────────────────────────────────────────────
    for item in spec.get("sections", []):
        stype = item.get("type", "paragraph")
        if stype in ("heading1", "heading2", "heading3"):
            add_paragraph(item, stype)
        elif stype == "paragraph":
            add_paragraph(item, "body")
        elif stype == "page_break":
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
        elif stype == "separator":
            p = doc.add_paragraph()
            run = p.add_run("—" * 20)
            run.font.size = Pt(10)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stype == "list":
            style = _resolve_style(doc_type, "body", item.get("font"), spec_styles)
            for i, it in enumerate(item.get("items", []), 1):
                p = doc.add_paragraph()
                _apply_paragraph_format(p.paragraph_format, style)
                run = p.add_run(f"{i}. {it}")
                _set_run_font(run, style.get("font_name"), style.get("size_pt"), style.get("bold"))
        elif stype == "table":
            rows = item.get("rows", [])
            if not rows:
                continue
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            style = _resolve_style(doc_type, "table", item.get("font"), spec_styles)
            for ri, row in enumerate(rows):
                for ci in range(ncols):
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    p = cell.paragraphs[0]
                    _apply_paragraph_format(p.paragraph_format, style)
                    run = p.add_run(row[ci] if ci < len(row) else "")
                    _set_run_font(run, style.get("font_name"), style.get("size_pt"), style.get("bold"))

    # 自动创建输出目录（不存在时）
    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    doc.save(output_path)
    return {"ok": True, "path": output_path}
